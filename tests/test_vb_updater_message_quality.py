"""A VB named by `pubRxStatus` carries a subscription's health, not a signal.

SEL Architect declares, on `<Private type="SEL_GooseSubscription">`, which bit
receives the health of each GOOSE subscription: it goes to 1 when the
publisher stops arriving. Nothing here could tell that apart from an ordinary
virtual bit, and it showed in two places.

Its `<ExtRef>` names a control block and no data attribute, so the Signal
column ran the full `IED/LD/LN/CB.ld.ln.do.da` template over the absent halves
and printed `QPC1_UPC2/ANN/LLN0/GoSB00....` -- four bare dots standing in for
a logical node, a data object and a data attribute that do not exist. And when
its `desc` was empty it was renamed `reserva`, the label for a slot foreseen
and unused, on a live GOOSE supervision bit: 5 of the 202 in
`samples/substation_demo.scd`.

The other 197 carry a real description ("FALHA GOOSE LT2 UPC1"), and those are
pinned here too: the marker never overwrites what the engineer wrote.
"""

from __future__ import annotations

from pacct.web.vb_updater import (
    _MESSAGE_QUALITY_LABEL,
    _RESERVA_LABEL,
    _new_comments_from_scd,
    _render_compare_page,
    build_vb_descriptions_xlsx,
    extract_vb_extref_rows_from_scd_ied,
    extract_vb_map_from_scd_ied,
)
from tests import gle_fixtures as fx

_IED = "QPC1_TR1"

# One IED with, in order: a data subscription (VB105), a health bit that has a
# description (VB051), a health bit that has none (VB052), and a plain VB the
# private block never mentions (VB060).
_SCD_BYTES = fx.scd(fx.ied(
    _IED,
    fx.extref("VB105", "TR1 UPC1 52A"),
    fx.rx_status_extref("VB051", "TR1 UPC1 FALHA GOOSE", cb="GoSB00"),
    fx.rx_status_extref("VB052", None, cb="GoSB01"),
    fx.extref("VB060", None),
    private=fx.goose_subscriptions(
        ("GoSB00", "VB051"),
        ("GoSB01", "VB052"),
        ("GoSB02", ""),      # a subscription whose health is not mapped
    ),
))


def _scd(tmp_path):
    p = tmp_path / "projeto.scd"
    p.write_bytes(_SCD_BYTES)
    return p


class TestWhichVbsAreMessageQuality:
    def test_only_the_declared_ones(self, tmp_path):
        vb_map = extract_vb_map_from_scd_ied(_scd(tmp_path), _IED)
        assert vb_map.quality == frozenset({"VB51", "VB52"})

    def test_the_descriptions_are_unchanged_by_the_new_reading(self, tmp_path):
        vb_map = extract_vb_map_from_scd_ied(_scd(tmp_path), _IED)
        assert vb_map.descs == {
            "VB105": "TR1 UPC1 52A",
            "VB51": "TR1 UPC1 FALHA GOOSE",
            "VB52": "",
            "VB60": "",
        }

    def test_an_ied_the_scd_does_not_have_is_empty_and_not_an_error(self, tmp_path):
        vb_map = extract_vb_map_from_scd_ied(_scd(tmp_path), "NAO_EXISTE")
        assert vb_map.descs == {} and vb_map.quality == frozenset()


class TestTheSignalColumn:
    def _rows(self, tmp_path):
        return {r["vb"]: r for r in
                extract_vb_extref_rows_from_scd_ied(_scd(tmp_path), _IED)}

    def test_a_health_bit_says_what_it_is_and_stops_at_the_control_block(self, tmp_path):
        row = self._rows(tmp_path)["VB51"]
        assert row["quality"] is True
        assert row["signal"] == f"{_MESSAGE_QUALITY_LABEL} — QPC1_UPC2/ANN/LLN0/GoSB00"

    def test_the_four_bare_dots_are_gone(self, tmp_path):
        # The old string was `QPC1_UPC2/ANN/LLN0/GoSB00....`: the template run
        # over a logical node, a data object and a data attribute that a
        # health bit does not have. Asserted against what the same attributes
        # still produce on the ordinary path, so this pins the difference and
        # not merely the absence of dots.
        from pacct.web.vb_updater import _EXTREF_FIELDS, _format_extref_signal

        attrs = dict.fromkeys(_EXTREF_FIELDS, "")
        attrs.update(iedName="QPC1_UPC2", srcLDInst="ANN",
                     srcLNClass="LLN0", srcCBName="GoSB00")
        assert _format_extref_signal(attrs) == "QPC1_UPC2/ANN/LLN0/GoSB00...."
        assert not self._rows(tmp_path)["VB51"]["signal"].endswith("....")

    def test_a_data_subscription_still_gets_the_full_signature(self, tmp_path):
        row = self._rows(tmp_path)["VB105"]
        assert row["quality"] is False
        assert row["signal"].startswith("QPC1_UPC2/")
        assert _MESSAGE_QUALITY_LABEL not in row["signal"]

    def test_the_spreadsheet_carries_the_same_string(self, tmp_path):
        from io import BytesIO

        from openpyxl import load_workbook

        blob = build_vb_descriptions_xlsx(scd_path=_scd(tmp_path), ied_names=[_IED])
        wb = load_workbook(BytesIO(blob))
        ws = wb[wb.sheetnames[0]]
        signals = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
                   for r in range(5, ws.max_row + 1)}
        assert signals["VB51"] == (
            f"{_MESSAGE_QUALITY_LABEL} — QPC1_UPC2/ANN/LLN0/GoSB00")


class TestTheCommentWrittenIntoTheGle:
    def _comments(self, tmp_path):
        comments, with_desc, reserva, quality = _new_comments_from_scd(
            _scd(tmp_path), _IED)
        return comments, {"with_desc": with_desc, "reserva": reserva,
                          "quality": quality}

    def test_an_empty_health_bit_is_no_longer_called_a_spare(self, tmp_path):
        comments, _ = self._comments(tmp_path)
        assert comments["VB52"] == _MESSAGE_QUALITY_LABEL

    def test_a_plain_empty_vb_is_still_a_spare(self, tmp_path):
        comments, _ = self._comments(tmp_path)
        assert comments["VB60"] == _RESERVA_LABEL

    def test_a_description_the_engineer_wrote_is_never_overwritten(self, tmp_path):
        comments, _ = self._comments(tmp_path)
        assert comments["VB51"] == "TR1 UPC1 FALHA GOOSE"

    def test_the_counts_say_what_they_mean(self, tmp_path):
        _, stats = self._comments(tmp_path)
        # `quality` counts what these VBs ARE (both of them), across the
        # branch that kept a description and the branch that was relabelled;
        # `reserva` is only VB60 now, where it used to be VB52 as well.
        assert stats == {"with_desc": 2, "reserva": 1, "quality": 2}


class TestTheComparePage:
    def _html(self, tmp_path):
        gle = tmp_path / "GL1.gle.xml"
        gle.write_bytes(fx.SAMPLE_GLE)
        return _render_compare_page(
            "TR1", _IED, "GL1", gle, _scd(tmp_path))

    def test_a_health_bit_row_is_tagged(self, tmp_path):
        html = self._html(tmp_path)
        assert f'<div class="tag">{_MESSAGE_QUALITY_LABEL}</div>' in html

    def test_the_tag_is_not_put_on_an_ordinary_vb(self, tmp_path):
        html = self._html(tmp_path)
        assert html.count(f'<div class="tag">{_MESSAGE_QUALITY_LABEL}</div>') == 2

    def test_the_summary_counts_them(self, tmp_path):
        assert f"2 {_MESSAGE_QUALITY_LABEL}" in self._html(tmp_path)
