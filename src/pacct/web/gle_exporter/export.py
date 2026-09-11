"""The two files this tool writes, and the one it reads back.

The spreadsheet is the round trip: `build_xlsx_for_selections` writes one sheet
per (relay, GLE) and `parse_xlsx_to_updates` reads that same sheet back after
the user has edited the Comment column. Both directions live here, together,
because the markers and the column numbers below are the format -- splitting
them across two modules would be one format written down twice.

`apply_xlsx_updates_to_rdb` is the other output: the edited comments applied to
the RDB, through `rdb_write` and nowhere else. It reads first and writes once,
for the reason its own docstring gives.

The spreadsheet the import route takes back is one this tool produced. It is
deliberately not a project input -- it never enters the library picker -- which
is why the format markers exist at all: they are how a sheet says it came from
here.
"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import Any

from sellib import rdb as rdb_loader
from sellib.rdb import RdbInfo

from pacct.web import rdb_write
from pacct.web.gle_exporter.model import (
    PortUpdates,
    diff_updates_against_gle,
    extract_port_instances_from_gle,
    update_port_comments_in_gle_bytes,
)
from pacct.web.rdb_write import (
    resolve_gle_stream_path as _resolve_gle_stream_path,
)
from pacct.web.xlsx_names import sanitize_sheet_name as _sanitize_sheet_name

_logger = logging.getLogger(__name__)


# Markers: they identify an xlsx as coming from this tool's export.
_XLSX_RELAY_MARKER = "Relay:"
_XLSX_GLE_MARKER = "GLE:"
# Row 4 header (each column = one cell in this header).
_XLSX_HEADERS = (
    "Page", "Element ID", "Type", "Variable", "Side", "Port", "Label", "Comment",
)
# 1-based column of each field (stable for the parser and the tests).
_COL_PAGE, _COL_EID, _COL_TYPE, _COL_VAR, _COL_SIDE, _COL_PORT, _COL_LABEL, _COL_CMT = range(1, 9)

# Output directory for the xlsx (reuses the /download sandbox).
# Uploads and outputs live in cache/sessions/<sid>/ (see pacct.web.session);
# there is no directory shared between users any more.


# -----------------------------------------------------------------------------
# Excel I/O
# -----------------------------------------------------------------------------

def build_xlsx_for_selections(
    selections: list[dict],
) -> bytes:
    """Build an .xlsx (bytes) with one sheet per selection.

    Each `selection` is a dict {relay, gle, gle_path, relay_model}. The sheet
    holds:
      A1: "Relay:"      B1: <relay_name>
      A2: "GLE:"        B2: <gle_name>
      Row 3: blank
      Row 4: headers [Page, Element ID, Type, Variable, Side, Port, Label, Comment]
      Row 5+: one row per port of each exportable element.

    `relay_model` (optional) enables the fixed port labels (S/R/Q/...);
    without it the Label column stays empty but the table is still valid.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    # A fresh `Workbook()` always has one sheet, but `active` is
    # Optional in the stubs -- every sheet here is added by name.
    default = wb.active
    if default is not None:
        wb.remove(default)

    used: set[str] = set()
    header_font = Font(bold=True, color="FFFFFFFF")
    header_fill = PatternFill("solid", fgColor="FF1F6FEB")
    meta_font = Font(bold=True)

    for sel in selections:
        relay = sel["relay"]
        gle = sel["gle"]
        gle_path = sel["gle_path"]
        relay_model = sel.get("relay_model")
        ports = extract_port_instances_from_gle(gle_path, relay_model=relay_model)

        sheet_base = f"{relay} {gle}".strip()
        ws = wb.create_sheet(title=_sanitize_sheet_name(sheet_base, used))
        ws["A1"] = _XLSX_RELAY_MARKER
        ws["B1"] = relay
        ws["A1"].font = meta_font
        ws["A2"] = _XLSX_GLE_MARKER
        ws["B2"] = gle
        ws["A2"].font = meta_font

        for col, val in enumerate(_XLSX_HEADERS, start=1):
            c = ws.cell(row=4, column=col, value=val)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="left")

        for i, p in enumerate(ports, start=5):
            ws.cell(row=i, column=_COL_PAGE,  value=p.page)
            ws.cell(row=i, column=_COL_EID,   value=p.element_id)
            ws.cell(row=i, column=_COL_TYPE,  value=p.xml_type)
            ws.cell(row=i, column=_COL_VAR,   value=p.name)
            ws.cell(row=i, column=_COL_SIDE,  value=p.side)
            ws.cell(row=i, column=_COL_PORT,  value=p.port_index)
            ws.cell(row=i, column=_COL_LABEL, value=p.label)
            ws.cell(row=i, column=_COL_CMT,   value=p.comment)

        # Widths: sized roughly so nothing shows up cut off.
        widths = {
            _COL_PAGE: 28, _COL_EID: 10, _COL_TYPE: 12, _COL_VAR: 18,
            _COL_SIDE: 8,  _COL_PORT: 6, _COL_LABEL: 10, _COL_CMT: 60,
        }
        for col_idx, w in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = w
        ws.freeze_panes = "A5"

    if not wb.sheetnames:
        wb.create_sheet(title="empty")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Type returned by the xlsx parser:
#   {(relay, gle) -> {element_id -> {(side, port_index) -> new_comment}}}
XlsxUpdates = dict[tuple[str, str], PortUpdates]


def parse_xlsx_to_updates(xlsx_bytes: bytes) -> XlsxUpdates:
    """Read an xlsx built by `build_xlsx_for_selections` (possibly edited)
    and return the port-by-port updates structure.

    Rules:
      - A sheet must have A1='Relay:' and A2='GLE:' (the markers). Other
        sheets are silently ignored.
      - Rows with an empty Element ID or an unknown Side are skipped.
      - If the Comment cell is None (cell erased/never filled in the
        original export), it does NOT enter updates -> the port is untouched.
      - An explicit "" Comment = a request to clear it (becomes <comment />).
      - Side is normalised to "input"/"output" (case insensitive).
    """
    from openpyxl import load_workbook

    wb = load_workbook(filename=BytesIO(xlsx_bytes), data_only=True, read_only=True)
    out: XlsxUpdates = {}
    try:
        for ws in wb.worksheets:
            a1 = ws.cell(row=1, column=1).value
            a2 = ws.cell(row=2, column=1).value
            if not (isinstance(a1, str) and a1.strip() == _XLSX_RELAY_MARKER):
                continue
            if not (isinstance(a2, str) and a2.strip() == _XLSX_GLE_MARKER):
                continue
            relay = ws.cell(row=1, column=2).value
            gle = ws.cell(row=2, column=2).value
            if not relay or not gle:
                continue
            relay = str(relay).strip()
            gle = str(gle).strip()
            per_gle: PortUpdates = out.setdefault((relay, gle), {})
            for row in ws.iter_rows(min_row=5, max_col=len(_XLSX_HEADERS),
                                     values_only=True):
                if row is None:
                    continue
                # Cells out of a spreadsheet the user uploaded: any of
                # openpyxl's value types, or None for a short row.
                # Every read below either `str()`s it or converts
                # inside a try, which is what makes that safe.
                cells: list[Any] = list(row) + [None] * len(_XLSX_HEADERS)
                _page = cells[_COL_PAGE - 1]
                eid = cells[_COL_EID - 1]
                _xml_type = cells[_COL_TYPE - 1]
                _var = cells[_COL_VAR - 1]
                side_raw = cells[_COL_SIDE - 1]
                port_raw = cells[_COL_PORT - 1]
                _label = cells[_COL_LABEL - 1]
                cmt = cells[_COL_CMT - 1]
                if eid is None:
                    continue
                eid_str = str(eid).strip()
                if not eid_str.isdigit():
                    continue
                side = str(side_raw or "").strip().lower()
                if side not in ("input", "output"):
                    continue
                try:
                    port_idx = int(port_raw) if port_raw is not None else 0
                except (TypeError, ValueError):
                    continue
                if cmt is None:
                    continue  # empty cell (not edited) -> do not touch
                entry = per_gle.setdefault(eid_str, {})
                entry[(side, port_idx)] = str(cmt).strip()
    finally:
        wb.close()
    return {k: v for k, v in out.items() if v}


# -----------------------------------------------------------------------------
# Orchestrator: applies the xlsx edits to the RDB and produces a new RDB
# -----------------------------------------------------------------------------

def apply_xlsx_updates_to_rdb(
    *,
    rdb_info: RdbInfo,
    xlsx_updates: XlsxUpdates,
    output_path: Path,
    job=None,
) -> dict:
    """Apply the xlsx changes to the RDB and produce `output_path`.

    `xlsx_updates`: (relay, gle) -> {element_id: {(side, port_idx): comment}}.

    Two passes, and the order matters. The FIRST one only reads: it resolves
    each (relay, GLE) to its stream in the OLE and computes the new bytes,
    without touching disk. Only if ALL of them succeed does the second pass
    write, in one go, via `rdb_write.write_streams` -- which uses
    `write_stream` when every stream keeps its size and rebuilds the
    container when one does not, always atomically onto the destination.

    It is all-or-nothing on purpose. Before, each selection was written into
    the output RDB as soon as it was ready: if the third failed, the first
    two were already in the file, the response still said `ok: True`, and
    that half-applied RDB went into the project library just like a complete
    one. A half-written settings file is indistinguishable from a whole one
    once it leaves here.

    Returns {ok, output_path, results, succeeded, failed, totals, method}.
    With `ok: False`, NO file was written and `output_path` does not exist.
    """
    from sellib.models import relay_models as _rm

    # RelayModel cache by name (resolved only once per relay).
    model_cache: dict[str, object] = {}
    def _model_for(relay_name: str):
        if relay_name in model_cache:
            return model_cache[relay_name]
        entry_relay = next(
            (r for r in rdb_info.relays if r.name == relay_name), None,
        )
        m = _rm.lookup(entry_relay.model) if (entry_relay and entry_relay.model) else None
        model_cache[relay_name] = m
        return m

    results: list[dict] = []
    totals = {"elements_touched": 0, "ports_updated": 0,
              "ports_skipped": 0, "elements_missing": 0}
    streams: dict[tuple[str, ...], bytes] = {}

    # -- pass 1: read only ----------------------------------------------------
    if job:
        job.stage("Conferindo os GLEs alterados", 10)
    for (relay, gle), elem_updates in xlsx_updates.items():
        # One row of `results`: heterogeneous by design -- the flag, the
        # stats dict, the byte count and the reason all live in it.
        entry: dict[str, Any] = {"relay": relay, "gle": gle}
        gle_entry = rdb_loader.find_gle(rdb_info, relay, gle)
        if gle_entry is None or not gle_entry.fs_path.is_file():
            entry["ok"] = False
            entry["error"] = "GLE nao encontrado no RDB"
            results.append(entry)
            continue

        # Filter out updates that already match the current content.
        delta = diff_updates_against_gle(
            gle_entry.fs_path, elem_updates,
            relay_model=_model_for(relay),
        )
        if not delta:
            entry["ok"] = True
            entry["stats"] = {"elements_touched": 0, "ports_updated": 0,
                              "ports_skipped": 0, "elements_missing": 0}
            entry["note"] = "nada a aplicar (valores ja batem)"
            results.append(entry)
            continue

        stream_parts = _resolve_gle_stream_path(rdb_info.extract_dir,
                                                gle_entry.fs_path)
        if not stream_parts:
            entry["ok"] = False
            entry["error"] = f"stream nao localizado no OLE: {gle_entry.rel_path}"
            results.append(entry)
            continue

        # The read comes from the EXTRACTED file, not from the OLE: it is
        # the same content, and this way pass 1 never opens the RDB.
        original = gle_entry.fs_path.read_bytes()
        updated, stats = update_port_comments_in_gle_bytes(original, delta)
        streams[tuple(stream_parts)] = updated
        entry["ok"] = True
        entry["stats"] = stats
        entry["original_stream_bytes"] = len(original)
        for k in totals:
            totals[k] += stats.get(k, 0)
        results.append(entry)

    succeeded = sum(1 for r in results if r.get("ok"))
    failed = len(results) - succeeded
    if failed:
        return {
            "ok": False,
            "error": (f"{failed} de {len(results)} seleção(ões) falharam; "
                      "nenhum RDB foi gravado."),
            "results": results,
            "succeeded": succeeded,
            "failed": failed,
            "totals": totals,
        }

    # -- pass 2: one write ---------------------------------------------------
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        if streams:
            method = rdb_write.write_streams(rdb_info.rdb_path, output_path,
                                             streams, job=job)
        else:
            # Every selection was a no-op. The user still gets the RDB
            # (which is what happened before, when the copy came first),
            # only with no stream rewritten.
            rdb_write.copy_only(rdb_info.rdb_path, output_path)
            method = "copy"
    except rdb_write.RdbWriteError as e:
        return {
            "ok": False,
            "error": str(e),
            "results": results,
            "succeeded": succeeded,
            "failed": failed,
            "totals": totals,
        }

    return {
        "ok": True,
        "output_path": str(output_path),
        "method": method,
        "results": results,
        "succeeded": succeeded,
        "failed": failed,
        "totals": totals,
    }

